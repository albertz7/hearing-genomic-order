from __future__ import annotations

import csv
import json
import pathlib

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = pathlib.Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output"
DOCX_DIR = OUTPUT / "docx"
SUMMARY = ROOT / "results" / "summary.json"
GENOMES = ROOT / "data" / "selected_genomes.csv"


REFERENCES = [
    "Wells JM, Raju BC, Hung HY, Weisburg WG, Mandelco-Paul L, Brenner DJ. Xylella fastidiosa gen. nov., sp. nov: Gram-negative, xylem-limited, fastidious plant bacteria related to Xanthomonas spp. International Journal of Systematic Bacteriology. 1987;37:136-143.",
    "Hopkins DL, Purcell AH. Xylella fastidiosa: cause of Pierce's disease of grapevine and other emergent diseases. Plant Disease. 2002;86:1056-1066.",
    "Simpson AJG, Reinach FC, Arruda P, Abreu FA, Acencio M, Alvarenga R, et al. The genome sequence of the plant pathogen Xylella fastidiosa. Nature. 2000;406:151-157.",
    "Bhattacharyya A, Stilwagen S, Ivanova N, D'Souza M, Bernal A, Lykidis A, et al. Whole-genome comparative analysis of three phytopathogenic Xylella fastidiosa strains. Proceedings of the National Academy of Sciences. 2002;99:12403-12408.",
    "Schaad NW, Postnikova E, Lacy G, Fatmi M, Chang CJ. Xylella fastidiosa subspecies: X. fastidiosa subsp. piercei, subsp. nov., X. fastidiosa subsp. multiplex subsp. nov., and X. fastidiosa subsp. pauca subsp. nov. Systematic and Applied Microbiology. 2004;27:290-300.",
    "Chen J, Xie G, Han S, Chertkov O, Sims D, Civerolo EL. Whole-genome sequences of two Xylella fastidiosa strains (M12 and M23) causing almond leaf scorch disease in California. Journal of Bacteriology. 2010;192:4534.",
    "Sicard A, Zeilinger AR, Vanhove M, Schartel TE, Beal DJ, Daugherty MP, Almeida RPP. Xylella fastidiosa: insights into an emerging plant pathogen. Annual Review of Phytopathology. 2018;56:181-202.",
    "Chatterjee S, Almeida RPP, Lindow S. Living in two worlds: the plant and insect lifestyles of Xylella fastidiosa. Annual Review of Phytopathology. 2008;46:243-271.",
    "Saponari M, Boscia D, Nigro F, Martelli GP. Identification of DNA sequences related to Xylella fastidiosa in oleander, almond and olive trees exhibiting leaf scorch symptoms in Apulia, southern Italy. Journal of Plant Pathology. 2013;95:668.",
    "Coletta-Filho HD, Castillo AI, Laranjeira FF, de Andrade EC, Silva NT, de Souza AA, Bossi ME, Lopes JRS, Almeida RPP. Citrus variegated chlorosis: an overview of 30 years of research and disease management. Tropical Plant Pathology. 2020;45:175-191.",
    "Nunney L, Vickerman DB, Bromley RE, Russell SA, Hartman JR, Morano LD, Stouthamer R. Recent evolutionary radiation and host plant specialization in the Xylella fastidiosa subspecies native to the United States. Applied and Environmental Microbiology. 2013;79:2189-2200.",
    "National Center for Biotechnology Information. NCBI Datasets Genome Data Package: Xylella fastidiosa complete RefSeq assemblies. Accessed July 12, 2026. https://www.ncbi.nlm.nih.gov/datasets/",
    "Sayers EW, Cavanaugh M, Clark K, Ostell J, Pruitt KD, Karsch-Mizrachi I. GenBank. Nucleic Acids Research. 2022;50:D161-D164.",
    "Temple MD. An auditory display tool for DNA sequence analysis. BMC Bioinformatics. 2017;18:221. doi:10.1186/s12859-017-1632-x.",
    "Martin Z, Meagher J, Barker D. Using sound to understand protein sequence data: new sonification algorithms for protein sequences and multiple sequence alignments. BMC Bioinformatics. 2021;22:456. doi:10.1186/s12859-021-04362-7.",
    "Plaisier H, Meagher J, Barker D. DNA sonification for public engagement in bioinformatics. BMC Research Notes. 2021;14:273. doi:10.1186/s13104-021-05685-7.",
    "Hermann T, Hunt A, Neuhoff JG, editors. The Sonification Handbook. Logos Verlag; 2011.",
    "Kramer G. Auditory Display: Sonification, Audification, and Auditory Interfaces. Addison-Wesley; 1994.",
    "Shannon CE. A mathematical theory of communication. Bell System Technical Journal. 1948;27:379-423, 623-656.",
    "Pérez F, Granger BE. IPython: a system for interactive scientific computing. Computing in Science and Engineering. 2007;9:21-29.",
    "Harris CR, Millman KJ, van der Walt SJ, Gommers R, Virtanen P, Cournapeau D, et al. Array programming with NumPy. Nature. 2020;585:357-362.",
    "Python Software Foundation. Python Language Reference, version 3.12. Accessed July 12, 2026. https://www.python.org/",
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def apply_base_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def add_title(doc: Document, title: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run(subtitle).italic = True


def h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_genome_table(doc: Document, rows: list[dict]) -> None:
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Accession", "Strain", "Host/source", "Location", "Length (bp)"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        set_cell_shading(cell, "EDEDED")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        values = [
            row["accession"],
            row["strain"],
            row["host_or_source"] or "not reported",
            row["geo_loc_name"] or "not reported",
            row["sequence_length"],
        ]
        for cell, value in zip(cells, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)


def build_manuscript() -> pathlib.Path:
    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    genomes = list(csv.DictReader(GENOMES.open(encoding="utf-8")))
    doc = Document()
    apply_base_styles(doc)
    add_title(
        doc,
        "Mapping Mutations to Music: Sonification and Comparative Analysis of Xylella fastidiosa Genomic Data",
        "Revised manuscript prepared for student-journal submission",
    )

    h(doc, "Abstract")
    p(
        doc,
        "Context: Sonification can make sequence patterns audible, but biological usefulness requires validation against known genomic differences rather than subjective listening alone. Objective: This study tested whether a deterministic DNA-to-MIDI encoding preserves measurable variation in Xylella fastidiosa, a bacterial plant pathogen associated with Pierce's disease, citrus variegated chlorosis, almond leaf scorch, and olive quick decline syndrome. Methods: Twenty complete RefSeq assemblies were downloaded from NCBI, including grapevine, citrus, coffee, plum, almond, olive, oak, and other host-associated strains. The first 180,000 unambiguous bases of each genome were encoded by a fixed codon-to-pitch map. Audio-derived features included pitch distribution, interval size, pitch entropy, and summary statistics. Validation used 190 pairwise genome comparisons, 300 synthetic mutation trials with known substitution rates from 0.05% to 2.0%, and shuffled-sequence null controls preserving nucleotide composition. Results: Pairwise audio-feature distances were moderately correlated with sequence-feature distances (Pearson r = 0.441). A holdout linear model predicted synthetic mutation rate from audio features with R2 = 0.916 and mean absolute error = 0.00127. Shuffled null sequences had a much larger median audio distance from originals (1.308) than synthetic mutation trials (0.0159). Conclusions: Sonification did not replace alignment or variant calling, but it produced reproducible audio features that tracked controlled sequence change. The method is best interpreted as an exploratory and educational display layer for genomic comparison, not as a diagnostic assay.",
    )
    p(doc, "Keywords: Xylella fastidiosa; genome sonification; auditory display; comparative genomics; synthetic mutation benchmark; MIDI.")

    h(doc, "Introduction")
    p(
        doc,
        "Xylella fastidiosa is not a virus; it is a Gram-negative, xylem-limited bacterium first described by Wells and colleagues [1]. It is transmitted by xylem-feeding insects and causes economically important plant diseases, including Pierce's disease of grapevine, citrus variegated chlorosis, almond leaf scorch, bacterial leaf scorch, and olive quick decline syndrome [2,7,8]. Its biology makes it an appealing case study for comparative genomics: complete genomes are small enough for student-scale analysis, but strains differ by host association, geography, and evolutionary history [3-7].",
    )
    p(
        doc,
        "The original version of this project attempted to compare only three accessions and described the resulting sounds qualitatively. That approach could not support the biological claims it made. Short or incomplete nucleotide records are not equivalent to complete genomes, subjective descriptions such as 'chaotic' or 'harmonic' are not measurements, and BLASTp cannot validate an audio file. The revised study therefore asks a narrower and testable question: when X. fastidiosa DNA is converted to musical events by a fixed rule, do quantitative audio features preserve measurable sequence differences?",
    )
    p(
        doc,
        "Prior studies show that sonification can support exploration and communication of DNA or protein patterns [14-16], and broader auditory-display research provides design principles for mapping data to sound [17,18]. However, a sonification method must be evaluated against ground truth. This study therefore combines real NCBI genomes with synthetic mutations and shuffled null controls. The hypothesis was that audio-feature distances would increase with sequence divergence and would distinguish low-rate synthetic mutation from randomized controls.",
    )

    h(doc, "Materials and Methods")
    h(doc, "Genome Selection", 2)
    p(
        doc,
        "Genome metadata were retrieved from the NCBI Datasets API on July 12, 2026. Assemblies were restricted to X. fastidiosa records labeled Complete Genome in RefSeq. Twenty assemblies were selected to include classic reference strains and a range of hosts and locations. Table 1 lists all accessions used. The FASTA files and selection metadata are stored locally in data/selected_genomes.csv and data/fastas/.",
    )
    p(doc, "Table 1. Complete RefSeq Xylella fastidiosa assemblies used in the revised analysis.")
    add_genome_table(doc, genomes)

    h(doc, "Preprocessing", 2)
    p(
        doc,
        "Each FASTA file was parsed by removing headers and retaining only A, C, G, and T. Ambiguous characters were excluded. To keep the analysis reproducible and computationally reasonable on a student workstation, the first 180,000 clean bases of each genome were used for the benchmark. This standardized segment length avoids giving longer assemblies more feature weight. The limitation is addressed in the Discussion.",
    )
    h(doc, "Sonification Mapping", 2)
    p(
        doc,
        "DNA was read in non-overlapping codons. The 64 possible codons were ordered lexicographically from AAA through TTT and mapped deterministically to MIDI pitches 36-83 using pitch = 36 + (codon index mod 48). Each event had the same duration in the preview MIDI files so that measured differences came from pitch content and pitch intervals rather than performance choices. This intentionally simple map was chosen because it is transparent, repeatable, and easy to audit.",
    )
    h(doc, "Audio and Sequence Features", 2)
    p(
        doc,
        "Sequence features included nucleotide frequencies, GC proportion, and 64 codon frequencies. Audio features included mean pitch, pitch standard deviation, median pitch, pitch entropy, mean absolute interval size, interval standard deviation, 10th and 90th pitch percentiles, and a 48-bin pitch histogram. Pairwise Euclidean distances were calculated for sequence features and audio features across the 20 genomes.",
    )
    h(doc, "Validation Design", 2)
    p(
        doc,
        "Three validation layers were used. First, all 190 genome pairs were compared to test whether audio-feature distance tracked sequence-feature distance. Second, synthetic point mutations were introduced at five known rates: 0.05%, 0.1%, 0.5%, 1.0%, and 2.0%. Three replicates per rate per genome produced 300 synthetic mutation records. Third, shuffled null sequences were generated by randomly permuting bases while preserving nucleotide composition. A 75/25 train-test split was used to test whether audio features could predict mutation rate in held-out synthetic examples. All random operations used seed 20260712.",
    )
    h(doc, "Reproducibility", 2)
    p(
        doc,
        "The pipeline is implemented in Python in analysis/select_and_download.py and analysis/run_benchmark.py. Outputs include results/pairwise_distances.csv, results/mutation_benchmark.csv, results/summary.json, and MIDI previews in output/audio/. The benchmark output hash was "
        + summary["feature_hash"]
        + ".",
    )

    h(doc, "Results")
    p(
        doc,
        f"The analysis used {summary['genomes']} complete RefSeq genomes and {summary['bases_per_genome_analyzed']:,} bases per genome. The 20 assemblies produced {summary['pairwise_comparisons']} pairwise comparisons. Audio-feature distance and sequence-feature distance were moderately correlated (Pearson r = {summary['pairwise_audio_sequence_distance_pearson_r']:.3f}). This result supports the hypothesis that the sonification map preserved some genomic similarity structure, but the correlation was not strong enough to claim that sound alone reconstructs phylogeny or identifies strain type.",
    )
    p(
        doc,
        f"In the controlled mutation benchmark, held-out mutation rate was predicted from audio-feature changes with R2 = {summary['mutation_rate_prediction_holdout_r2']:.3f} and mean absolute error = {summary['mutation_rate_prediction_holdout_mae']:.5f}. The median audio distance for synthetic mutation trials was {summary['median_synthetic_mutation_audio_distance']:.4f}, whereas the median audio distance for shuffled null controls was {summary['median_shuffled_null_audio_distance']:.3f}. The difference indicates that the audio representation responded to structured sequence changes while treating randomized sequences as far more distant.",
    )
    p(
        doc,
        "The MIDI previews confirmed that the pipeline produces playable sonifications, but the statistical analysis did not depend on human ratings of pleasantness, harmony, or perceived disorder. Listening can help a researcher notice patterns, but the evidence in this study comes from computed features and controlled comparisons.",
    )

    h(doc, "Discussion")
    p(
        doc,
        "The revised results support a limited claim: deterministic sonification can act as a reproducible display layer that preserves measurable aspects of sequence variation. The study does not show that sonification can diagnose host range, pathogenicity, or subspecies. Those biological conclusions require curated labels, alignments, gene annotation, variant calling, and independent experimental evidence. This distinction is important because X. fastidiosa disease outcome depends on host, vector, environment, and bacterial genotype [2,7,8].",
    )
    p(
        doc,
        "The strongest part of the revised design is the ground-truth mutation test. Synthetic substitutions provided known rates that could be compared against audio-feature changes. Shuffled controls provided a negative-control condition that preserved base composition while destroying local order. Together, these controls address the main weakness of the original manuscript: it described sounds but did not test whether the sounds corresponded to known sequence differences.",
    )
    p(
        doc,
        "Several limitations remain. First, the codon-to-pitch map is deterministic but still arbitrary; other musical encodings could emphasize amino acid class, codon degeneracy, GC content, motif recurrence, or gene boundaries. Second, the first 180,000 bases were analyzed for speed, so the current benchmark is a standardized segment comparison rather than a full-genome alignment. Third, Euclidean distance on summary features is simpler than whole-genome SNP analysis and should not be interpreted as evolutionary distance. Fourth, the dataset was not used to classify disease severity because host labels and pathogenic outcomes are not equivalent.",
    )
    p(
        doc,
        "Future work should compare multiple sonification maps against the same mutation benchmark, include complete genome alignments, use curated subspecies and host-pathogenicity labels, and test whether trained listeners can detect specific sequence changes under blinded conditions. A useful next benchmark would measure sensitivity and specificity for predefined mutations inserted at known positions.",
    )

    h(doc, "Conclusion")
    p(
        doc,
        "This project was rebuilt from a qualitative demonstration into a reproducible validation study. Across 20 complete X. fastidiosa RefSeq genomes, a transparent codon-to-MIDI mapping produced audio features that moderately tracked sequence-feature distances and strongly reflected controlled synthetic mutation rates. Sonification should not be presented as a replacement for standard bioinformatics. Its value is as an interpretable, educational, and exploratory companion to quantitative sequence analysis.",
    )

    h(doc, "Data and Code Availability")
    p(
        doc,
        "All scripts, selected accession metadata, benchmark outputs, and MIDI preview files are contained in the local project workspace. The main reproducible commands are: python analysis/select_and_download.py and python analysis/run_benchmark.py.",
    )

    h(doc, "References")
    for idx, ref in enumerate(REFERENCES, start=1):
        doc.add_paragraph(f"{idx}. {ref}")

    path = DOCX_DIR / "Mapping_Mutations_to_Music_Revised_Manuscript.docx"
    doc.save(path)
    return path


def build_response() -> pathlib.Path:
    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    apply_base_styles(doc)
    add_title(doc, "Response to NHSJS Reviewer Feedback")
    p(doc, "Dear Editor and Reviewers,")
    p(
        doc,
        "Thank you for the detailed review. I treated the decision letter as a roadmap for a full rebuild rather than a light revision. The manuscript now uses correct bacterial terminology, explicit RefSeq accessions, a larger complete-genome dataset, a deterministic and documented sonification map, synthetic mutation ground truth, shuffled negative controls, and quantitative statistics. The claims have been narrowed to what the data support.",
    )
    h(doc, "Major Changes")
    add_bullets(
        doc,
        [
            "Replaced the original three-record comparison with 20 complete RefSeq Xylella fastidiosa genomes.",
            "Removed unsupported viral terminology and corrected the organism biology throughout.",
            "Replaced qualitative sound descriptions with pairwise feature distances, synthetic mutation-rate prediction, and shuffled null controls.",
            "Removed the incorrect claim that BLASTp validates audio output.",
            "Added a transparent codon-to-MIDI mapping, preprocessing rules, scripts, metadata files, and reproducible output tables.",
            "Expanded and modernized the reference list, including the sonification papers requested by the reviewer.",
        ]
    )
    h(doc, "Point-by-Point Response")
    responses = [
        ("Aim and hypothesis unclear", "The Introduction now states a narrow hypothesis: audio-feature distance should increase with sequence divergence and distinguish synthetic mutations from shuffled null controls."),
        ("Organism described incorrectly", "The manuscript now identifies Xylella fastidiosa as a Gram-negative, xylem-limited bacterium and removes viral language."),
        ("Abstract too vague", "The abstract is now structured with context, objective, methods, results, and conclusion, including sample size and numerical findings."),
        ("Unsupported disease/accession claims", "Disease claims are tied to the Xylella literature, and Table 1 lists all accessions, strains, hosts or sources, locations, and sequence lengths."),
        ("Sample size inadequate", "The analysis now uses 20 complete RefSeq genomes plus 300 synthetic mutation records and shuffled null controls."),
        ("Incomplete records treated as genomes", "The selected dataset is restricted to RefSeq assemblies labeled Complete Genome. The problematic short accessions from the original draft are no longer used as genomes."),
        ("Reference genome undefined", "The revised design does not use an undefined 'minimally mutated' reference. Each genome serves as its own baseline for synthetic mutation testing."),
        ("Missing sonification literature", "Temple 2017, Martin 2021, and Plaisier 2021 were added and discussed."),
        ("Auditory value not justified", "The Discussion now frames sonification as an exploratory and educational display layer, not a replacement for alignment or variant calling."),
        ("Preprocessing not reproducible", "The Methods now specify FASTA parsing, base filtering, standardized 180,000-base segments, software files, output paths, and random seed."),
        ("Mapping not specified", "The Methods now define the codon ordering, MIDI pitch formula, features, and preview MIDI generation."),
        ("Transcription rationale unclear", "The revised pipeline no longer claims biological transcription to mRNA; it analyzes DNA codons directly."),
        ("BLASTp validation inappropriate", "This claim was removed. Validation now uses known synthetic mutation rates and negative controls."),
        ("No negative controls", "Shuffled-sequence null controls preserving nucleotide composition were added."),
        ("No quantitative acoustic features", "Pitch distribution, interval statistics, entropy, and pitch histograms are now measured."),
        ("No statistical tests", "The Results now report pairwise correlation, holdout R2, mean absolute error, and median distances for mutation and null conditions."),
        ("Overclaims", "Claims about pathogenicity, virulence detection, and diagnostic use were removed or explicitly limited."),
        ("Reference list too short", "The manuscript now contains 22 references, including primary Xylella genome/pathology sources and sonification/auditory-display sources."),
    ]
    for i, (issue, response) in enumerate(responses, start=1):
        p(doc, f"{i}. Reviewer concern: {issue}. Response: {response}")
    p(doc, "Sincerely,")
    p(doc, "Albert Zhang")
    path = DOCX_DIR / "NHSJS_Response_to_Reviewers.docx"
    doc.save(path)
    return path


def main() -> None:
    manuscript = build_manuscript()
    response = build_response()
    print(manuscript)
    print(response)


if __name__ == "__main__":
    main()
