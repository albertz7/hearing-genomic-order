from __future__ import annotations

import csv
import json
import pathlib
from collections import Counter, defaultdict

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = pathlib.Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output"
DOCX_DIR = OUTPUT / "docx"
PDF_DIR = OUTPUT / "pdf"
FIG_DIR = ROOT / "results" / "figures"


TITLE = "Hearing Genomic Order: Windowed Sonification for Representing Mutation and Anomaly Patterns in Xylella fastidiosa Assemblies"


REFERENCES = [
    "Wells JM, Raju BC, Hung HY, Weisburg WG, Mandelco-Paul L, Brenner DJ. Xylella fastidiosa gen. nov., sp. nov: Gram-negative, xylem-limited, fastidious plant bacteria related to Xanthomonas spp. International Journal of Systematic Bacteriology. 1987;37:136-143.",
    "Hopkins DL, Purcell AH. Xylella fastidiosa: cause of Pierce's disease of grapevine and other emergent diseases. Plant Disease. 2002;86:1056-1066.",
    "Simpson AJG, Reinach FC, Arruda P, Abreu FA, Acencio M, Alvarenga R, et al. The genome sequence of the plant pathogen Xylella fastidiosa. Nature. 2000;406:151-157.",
    "Bhattacharyya A, Stilwagen S, Ivanova N, D'Souza M, Bernal A, Lykidis A, et al. Whole-genome comparative analysis of three phytopathogenic Xylella fastidiosa strains. Proceedings of the National Academy of Sciences. 2002;99:12403-12408.",
    "Schaad NW, Postnikova E, Lacy G, Fatmi M, Chang CJ. Xylella fastidiosa subspecies: X. fastidiosa subsp. piercei, subsp. nov., X. fastidiosa subsp. multiplex subsp. nov., and X. fastidiosa subsp. pauca subsp. nov. Systematic and Applied Microbiology. 2004;27:290-300.",
    "Chen J, Xie G, Han S, Chertkov O, Sims D, Civerolo EL. Whole-genome sequences of two Xylella fastidiosa strains (M12 and M23) causing almond leaf scorch disease in California. Journal of Bacteriology. 2010;192:4534.",
    "Chatterjee S, Almeida RPP, Lindow S. Living in two worlds: the plant and insect lifestyles of Xylella fastidiosa. Annual Review of Phytopathology. 2008;46:243-271.",
    "Sicard A, Zeilinger AR, Vanhove M, Schartel TE, Beal DJ, Daugherty MP, Almeida RPP. Xylella fastidiosa: insights into an emerging plant pathogen. Annual Review of Phytopathology. 2018;56:181-202.",
    "Saponari M, Boscia D, Nigro F, Martelli GP. Identification of DNA sequences related to Xylella fastidiosa in oleander, almond and olive trees exhibiting leaf scorch symptoms in Apulia, southern Italy. Journal of Plant Pathology. 2013;95:668.",
    "Coletta-Filho HD, Castillo AI, Laranjeira FF, de Andrade EC, Silva NT, de Souza AA, Bossi ME, Lopes JRS, Almeida RPP. Citrus variegated chlorosis: an overview of 30 years of research and disease management. Tropical Plant Pathology. 2020;45:175-191.",
    "Nunney L, Vickerman DB, Bromley RE, Russell SA, Hartman JR, Morano LD, Stouthamer R. Recent evolutionary radiation and host plant specialization in the Xylella fastidiosa subspecies native to the United States. Applied and Environmental Microbiology. 2013;79:2189-2200.",
    "National Center for Biotechnology Information. NCBI Datasets Genome Data Package: Xylella fastidiosa RefSeq assemblies. Accessed July 12, 2026. https://www.ncbi.nlm.nih.gov/datasets/",
    "Sayers EW, Cavanaugh M, Clark K, Ostell J, Pruitt KD, Karsch-Mizrachi I. GenBank. Nucleic Acids Research. 2022;50:D161-D164.",
    "Temple MD. An auditory display tool for DNA sequence analysis. BMC Bioinformatics. 2017;18:221. doi:10.1186/s12859-017-1632-x.",
    "Martin Z, Meagher J, Barker D. Using sound to understand protein sequence data: new sonification algorithms for protein sequences and multiple sequence alignments. BMC Bioinformatics. 2021;22:456. doi:10.1186/s12859-021-04362-7.",
    "Plaisier H, Meagher J, Barker D. DNA sonification for public engagement in bioinformatics. BMC Research Notes. 2021;14:273. doi:10.1186/s13104-021-05685-7.",
    "Hermann T, Hunt A, Neuhoff JG, editors. The Sonification Handbook. Logos Verlag; 2011.",
    "Kramer G. Auditory Display: Sonification, Audification, and Auditory Interfaces. Addison-Wesley; 1994.",
    "Shannon CE. A mathematical theory of communication. Bell System Technical Journal. 1948;27:379-423, 623-656.",
    "Harris CR, Millman KJ, van der Walt SJ, Gommers R, Virtanen P, Cournapeau D, et al. Array programming with NumPy. Nature. 2020;585:357-362.",
    "Van Rossum G, Drake FL. Python 3 Reference Manual. Python Software Foundation; 2009.",
]


def font(size: int = 12, bold: bool = False):
    try:
        return ImageFont.truetype("arial.ttf", size=size)
    except OSError:
        return ImageFont.load_default()


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def apply_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    for style_name in ["Normal", "List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_after = Pt(6)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)


def h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_title(doc: Document) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(TITLE)
    set_run_font(run, size=14, bold=True)
    para.paragraph_format.space_after = Pt(8)


def shade_cell(cell, fill: str = "EDEDED") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_table(table, font_size: int = 9) -> None:
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    set_run_font(run, size=font_size)


def add_dataset_summary_table(doc: Document, genomes: list[dict]) -> None:
    counts = defaultdict(Counter)
    for row in genomes:
        counts[row["host_group"]][row["assembly_level"]] += 1
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Host group", "Complete", "Chromosome", "Scaffold", "Total"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
        shade_cell(table.rows[0].cells[i])
    repeat_header(table.rows[0])
    for group in ["grapevine", "almond", "olive", "coffee", "other_prunus"]:
        row = table.add_row().cells
        row[0].text = group
        row[1].text = str(counts[group]["Complete Genome"])
        row[2].text = str(counts[group]["Chromosome"])
        row[3].text = str(counts[group]["Scaffold"])
        row[4].text = str(sum(counts[group].values()))
    format_table(table, 9)


def add_accession_table(doc: Document, genomes: list[dict]) -> None:
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Group", "Accession", "Strain", "Assembly", "Host/source"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
        shade_cell(table.rows[0].cells[i])
    repeat_header(table.rows[0])
    for item in genomes:
        row = table.add_row().cells
        row[0].text = item["host_group"]
        row[1].text = item["accession"]
        row[2].text = item["strain"]
        row[3].text = item["assembly_level"]
        row[4].text = item["host_or_source"]
    format_table(table, 7)


def add_result_table(doc: Document, rows: list[dict]) -> None:
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Feature set", "Anomaly", "Mean AUC", "95% CI", "Precision@4"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
        shade_cell(table.rows[0].cells[i])
    repeat_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        cells[0].text = row["feature_kind"]
        cells[1].text = row["anomaly_type"].replace("_", " ")
        cells[2].text = f"{float(row['mean_auc']):.3f}"
        cells[3].text = f"{float(row['auc_ci_low']):.3f}-{float(row['auc_ci_high']):.3f}"
        cells[4].text = f"{float(row['mean_precision_at_k']):.3f}"
    format_table(table, 9)


def add_figure(doc: Document, path: pathlib.Path, caption: str, width: float = 5.9) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        set_run_font(run, size=10, italic=True)


def draw_bar_chart(path: pathlib.Path, title: str, labels: list[str], values: list[float], x_label: str = "") -> None:
    img = Image.new("RGB", (1200, 700), "white")
    draw = ImageDraw.Draw(img)
    title_font = font(30)
    label_font = font(22)
    small_font = font(18)
    draw.text((600, 35), title, anchor="mm", fill="black", font=title_font)
    left, top, width, height = 270, 110, 820, 470
    max_v = max(values + [1.0])
    bar_h = min(52, int(height / max(len(labels), 1)) - 12)
    for i, (label, value) in enumerate(zip(labels, values)):
        y = top + i * int(height / max(len(labels), 1)) + 8
        w = int(width * value / max_v)
        color = "#4C78A8" if "composition" not in label else "#A0A0A0"
        draw.text((left - 12, y + bar_h / 2), label, anchor="rm", fill="black", font=small_font)
        draw.rectangle((left, y, left + w, y + bar_h), fill=color)
        draw.text((left + w + 10, y + bar_h / 2), f"{value:.3f}", anchor="lm", fill="black", font=small_font)
    draw.line((left, top + height + 20, left + width, top + height + 20), fill="black", width=2)
    if x_label:
        draw.text((left + width / 2, top + height + 58), x_label, anchor="mm", fill="black", font=label_font)
    img.save(path)


def draw_grouped_auc_chart(path: pathlib.Path, rows: list[dict]) -> None:
    order = ["codon_order_shuffle", "clustered_substitution", "foreign_segment"]
    kinds = ["composition", "codon64", "amino_class"]
    colors = {"composition": "#A0A0A0", "codon64": "#4C78A8", "amino_class": "#F58518"}
    values = {(r["anomaly_type"], r["feature_kind"]): float(r["mean_auc"]) for r in rows}
    img = Image.new("RGB", (1300, 760), "white")
    draw = ImageDraw.Draw(img)
    draw.text((650, 42), "Reference-guided anomaly detection", anchor="mm", fill="black", font=font(32))
    left, top, plot_w, plot_h = 120, 120, 1080, 470
    draw.line((left, top + plot_h, left + plot_w, top + plot_h), fill="black", width=2)
    draw.line((left, top, left, top + plot_h), fill="black", width=2)
    for tick in [0, 0.25, 0.5, 0.75, 1.0]:
        y = top + plot_h - tick * plot_h
        draw.line((left - 5, y, left + plot_w, y), fill="#E6E6E6" if tick else "black", width=1)
        draw.text((left - 14, y), f"{tick:.2f}", anchor="rm", fill="black", font=font(16))
    group_w = plot_w / len(order)
    bar_w = 70
    for gi, anomaly in enumerate(order):
        center = left + group_w * gi + group_w / 2
        for ki, kind in enumerate(kinds):
            val = values[(anomaly, kind)]
            x = center + (ki - 1) * (bar_w + 10)
            y = top + plot_h - val * plot_h
            draw.rectangle((x - bar_w / 2, y, x + bar_w / 2, top + plot_h), fill=colors[kind])
            draw.text((x, y - 8), f"{val:.2f}", anchor="mb", fill="black", font=font(15))
        draw.text((center, top + plot_h + 42), anomaly.replace("_", " "), anchor="mm", fill="black", font=font(18))
    for i, kind in enumerate(kinds):
        x = 395 + i * 180
        draw.rectangle((x, 655, x + 24, 679), fill=colors[kind])
        draw.text((x + 34, 667), kind, anchor="lm", fill="black", font=font(18))
    draw.text((650, 720), "AUC; higher values mean true anomaly windows rank above background-mutated windows", anchor="mm", fill="black", font=font(18))
    img.save(path)


def draw_track(path: pathlib.Path, rows: list[dict]) -> None:
    scores = [float(r["score"]) for r in rows]
    labels = [int(r["label"]) for r in rows]
    img = Image.new("RGB", (1300, 520), "white")
    draw = ImageDraw.Draw(img)
    draw.text((650, 38), "Example sonified anomaly score track", anchor="mm", fill="black", font=font(30))
    left, top, plot_w, plot_h = 80, 90, 1120, 320
    draw.rectangle((left, top, left + plot_w, top + plot_h), outline="black", width=2)
    hi = max(scores) if scores else 1
    pts = []
    for i, score in enumerate(scores):
        x = left + plot_w * i / max(len(scores) - 1, 1)
        y = top + plot_h - plot_h * score / hi
        pts.append((x, y))
        if labels[i]:
            draw.rectangle((x - 5, top, x + 5, top + plot_h), fill="#FAD7A0")
    if len(pts) > 1:
        draw.line(pts, fill="#4C78A8", width=4)
    draw.text((650, 455), "Genome windows; shaded bands are planted anomaly windows and the blue line is codon64 audio-feature distance", anchor="mm", fill="black", font=font(18))
    img.save(path)


def make_png_figures(summary: dict, anomaly_rows: list[dict], detail_rows: list[dict]) -> dict[str, pathlib.Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    class_rows = summary["classification"]
    class_png = FIG_DIR / "classification_accuracy.png"
    draw_bar_chart(
        class_png,
        "Host-group representation test",
        [row["feature_kind"] for row in class_rows],
        [float(row["accuracy"]) for row in class_rows],
        "Leave-one-genome-out accuracy; chance = 0.20",
    )
    auc_png = FIG_DIR / "anomaly_auc.png"
    draw_grouped_auc_chart(auc_png, anomaly_rows)
    track_png = FIG_DIR / "example_anomaly_track.png"
    draw_track(track_png, detail_rows)
    return {"classification": class_png, "auc": auc_png, "track": track_png}


def load_inputs():
    summary = json.loads((ROOT / "results" / "summary.json").read_text(encoding="utf-8"))
    genomes = list(csv.DictReader((ROOT / "data" / "study_genomes.csv").open(encoding="utf-8")))
    anomaly = list(csv.DictReader((ROOT / "results" / "anomaly_detection_summary.csv").open(encoding="utf-8")))
    detail = list(csv.DictReader((ROOT / "results" / "example_anomaly_window_scores.csv").open(encoding="utf-8")))
    return summary, genomes, anomaly, detail


def build_manuscript() -> pathlib.Path:
    summary, genomes, anomaly, detail = load_inputs()
    figs = make_png_figures(summary, anomaly, detail)
    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    apply_styles(doc)
    add_title(doc)

    h(doc, "Abstract")
    p(doc, "Background/Objective: Genomic sonification can turn DNA sequence patterns into sound, but many demonstrations do not test whether sound-based features reveal information beyond ordinary sequence summaries. This study tested whether windowed sonification can represent local mutation and anomaly patterns in Xylella fastidiosa, a bacterial plant pathogen. Methods: Fifty NCBI RefSeq assemblies were selected from five host-associated groups: grapevine, almond, olive, coffee, and other Prunus hosts. Genomes were divided into 3,000 bp windows. Two deterministic sonification maps were evaluated: a 64-codon pitch map and an amino-acid-class pitch map. Audio features were compared with standard base/codon-composition features, and MIDI anomaly tracks were exported. The study tested host-group representation with leave-one-genome-out nearest-centroid classification and tested reference-guided anomaly detection using background substitutions plus planted clustered substitutions, foreign-segment replacements, and codon-order shuffles. Results: Codon64 sonification matched composition features for host-group classification (40% accuracy, p = 0.0196; chance = 20%). For codon-order shuffles, composition features failed (AUC = 0.499), while codon64 and amino-class sonification performed strongly (AUC = 0.974 and 0.958). For foreign-segment anomalies, codon64 sonification reached AUC = 0.993. Conclusions: Sonification did not replace alignment or variant calling. Its strongest use was as an interpretable representation of genomic order anomalies, especially changes that preserve composition but alter sequential structure.")
    p(doc, "Keywords: Xylella fastidiosa; genome sonification; anomaly detection; codon order; auditory display; comparative genomics; MIDI.")

    h(doc, "Introduction")
    p(doc, "Xylella fastidiosa is a Gram-negative, xylem-limited bacterial plant pathogen associated with Pierce's disease of grapevine, citrus variegated chlorosis, almond leaf scorch, coffee leaf scorch, and olive quick decline syndrome [1,2,7-10]. Its small bacterial genome, broad host range, and large public sequence record make it a useful system for student-scale comparative genomics [3-6,11-13].")
    p(doc, "The central problem addressed here is not whether music can replace standard bioinformatics. It cannot. Alignment, variant calling, gene annotation, and phylogenetic analysis remain the correct tools for biological inference. The open question is narrower: can sonification provide a useful representation layer for genomic changes that are difficult to notice from composition summaries alone?")
    p(doc, "Previous DNA and protein sonification studies have shown that biological sequences can be mapped to sound for exploration, education, and pattern display [14-18]. However, many sonification projects remain qualitative. They describe sequences as more or less musical without testing whether audio-derived features recover known sequence differences. A strong sonification study therefore needs explicit mappings, negative controls, baseline methods, and ground-truth anomalies.")
    p(doc, "This study reframes the project as a windowed anomaly-representation benchmark. The hypothesis was that deterministic sonification features would preserve some host-associated sequence signal and would be especially useful for detecting order-based anomalies, such as codon-order shuffles, that leave base and codon composition largely unchanged.")

    h(doc, "Materials and Methods")
    h(doc, "Study Design", 2)
    p(doc, "This was a computational methods study using public bacterial genome assemblies. The analysis had two parts. First, genome-level audio and composition features were tested for host-group representation. Second, window-level features were tested for reference-guided anomaly detection. The second task was the main experiment because it directly tested whether sonification can represent mutation or anomaly windows.")
    h(doc, "Genome Selection", 2)
    p(doc, "Genome metadata were retrieved from the NCBI Datasets API on July 12, 2026. Assemblies were restricted to RefSeq records for X. fastidiosa with assembly levels Complete Genome, Chromosome, or Scaffold. Records were quality-filtered using CheckM metadata when available: completeness had to be at least 92.5%, and contamination had to be no more than 3%. The final dataset contained 50 assemblies, with 10 assemblies in each host-associated group. Table 1 summarizes the dataset; the full accession list is included in Appendix Table A1.")
    p(doc, "Table 1. Study dataset summary by host-associated group and assembly level.")
    add_dataset_summary_table(doc, genomes)
    h(doc, "Preprocessing", 2)
    p(doc, "FASTA files were parsed by removing headers and retaining only A, C, G, and T. Each genome was divided into non-overlapping 3,000 bp windows. To keep the analysis balanced and reproducible, the first 60 complete windows from each assembly were used, giving 3,000 genome windows across the 50-assembly dataset.")
    h(doc, "Sonification Mappings", 2)
    p(doc, "Two deterministic sequence-to-pitch maps were evaluated. In the codon64 map, the 64 possible codons were ordered lexicographically from AAA through TTT and mapped to MIDI pitches 36 through 99. In the amino-class map, each codon was translated using the standard bacterial genetic code and assigned to one of seven biochemical pitch classes: hydrophobic, polar, acidic, basic, aromatic, special, or stop.")
    p(doc, "Audio features were calculated from the pitch sequence for each 3,000 bp window. Let X = (x1, x2, ..., xm) be the MIDI pitch sequence for a window. Pitch distribution was the normalized histogram of pitch values. Pitch entropy was calculated as H = -Σi P(xi) log2 P(xi), where P(xi) is the probability of pitch class i in the window and zero-probability terms were omitted [19]. Mean pitch was mean(X), and pitch spread was the standard deviation of X. First-order intervals were defined as d(j) = x(j+1) - x(j). Interval size was mean(|d(j)|), interval variation was sd(d(j)), and second-order interval change was mean(|d(j+1) - d(j)|). Distances between a reference window and a modified window were Euclidean distances between their feature vectors.")
    h(doc, "Baseline Features", 2)
    p(doc, "The baseline was a non-audio sequence summary containing base frequencies, GC proportion, and 64 codon frequencies. This baseline was intentionally strong for composition-changing anomalies. A meaningful sonification result therefore required audio features to perform well where composition features did not.")
    h(doc, "Host-Group Representation Test", 2)
    p(doc, "For each genome, window-level features were summarized by feature means and standard deviations. A leave-one-genome-out nearest-centroid classifier predicted the host-associated group for the held-out assembly. Label permutation testing with 50 permutations estimated whether accuracy exceeded chance. Because the dataset had five balanced groups, chance accuracy was 20%.")
    h(doc, "Reference-Guided Anomaly Test", 2)
    p(doc, "The main anomaly experiment simulated a practical usecase: a sample sequence is compared with a reference sequence, and sonification is used to make high-change windows audible or visually trackable. Each window first received low background noise through 0.1% random substitutions. Four anomaly windows per genome were then planted for each anomaly type. Clustered substitutions added 1.5% substitutions within selected windows. Foreign-segment anomalies replaced the middle fifth of a window with sequence from another host group. Codon-order shuffles randomly permuted codons across the middle 60% of a window, preserving codon composition while disrupting local order.")
    p(doc, "For each modified genome, the analysis ranked windows by distance from the original reference window. Area under the receiver operating characteristic curve (AUC) measured whether planted anomaly windows ranked above background-mutated windows. Precision@4 measured the fraction of true anomaly windows among the four highest-scoring windows. The value 4 was chosen because exactly four anomaly windows were planted per genome for each anomaly type, so Precision@4 directly measured whether the method placed the planted anomalies at the top of the ranked window list. Bootstrap confidence intervals used 500 resamples across genome-level AUC values.")
    h(doc, "Ethical Considerations and Reproducibility", 2)
    p(doc, "The study used only public bacterial genome assemblies and did not involve human participants, animals, private data, or clinical decision-making. Scripts were written in Python with NumPy support [20,21] and are provided in analysis/select_and_download.py, analysis/sonification_core.py, and analysis/run_benchmark.py. Outputs are stored in data/study_genomes.csv, results/classification_results.csv, results/anomaly_detection_summary.csv, results/summary.json, and output/audio/. The random seed was 20260712.")

    h(doc, "Results")
    h(doc, "Dataset", 2)
    p(doc, f"The final dataset contained {summary['study_genomes']} assemblies and five balanced host-associated groups: grapevine, almond, olive, coffee, and other Prunus, with 10 assemblies per group. The windowed analysis used {summary['max_windows_per_genome']} windows per assembly, producing 3,000 genome windows.")
    h(doc, "Host-Group Representation", 2)
    p(doc, "The host-group representation test was modest but above chance for codon64 sonification. Composition features and codon64 audio features each reached 40% leave-one-genome-out accuracy with permutation p = 0.0196. Amino-class sonification reached 30% accuracy with p = 0.0588. These results suggest that the codon-level audio representation retained some host-associated sequence signal, but the result should not be interpreted as a diagnostic classifier.")
    add_figure(doc, figs["classification"], "Figure 1. Leave-one-genome-out host-group classification. The codon64 audio representation matched the composition baseline, while the amino-class map was weaker.")
    h(doc, "Reference-Guided Anomaly Detection", 2)
    p(doc, "The anomaly benchmark showed where sonification was most useful. For clustered substitutions and foreign-segment replacements, the composition baseline was very strong because these changes altered base or codon frequencies. Codon64 and amino-class sonification also performed well, but they did not outperform composition for those composition-changing anomalies.")
    p(doc, "The clearest sonification-specific result appeared in the codon-order shuffle condition. Composition features performed at chance (mean AUC = 0.499, 95% CI 0.456-0.539) because codon counts were preserved. In contrast, codon64 sonification reached mean AUC = 0.974 (95% CI 0.967-0.980), and amino-class sonification reached mean AUC = 0.958 (95% CI 0.945-0.969). The interval structure of the audio representation captured sequence-order disruption that composition summaries missed.")
    p(doc, "Precision@4 showed the same pattern in a more practical ranked-list form. For codon-order shuffles, codon64 sonification recovered true anomaly windows among the top four scoring windows much more often than the composition baseline (Precision@4 = 0.670 versus 0.050), while amino-class sonification reached Precision@4 = 0.730. This supports the interpretation that sonification is most useful when the abnormal signal is in sequence order rather than simple codon abundance.")
    add_result_table(doc, anomaly)
    add_figure(doc, figs["auc"], "Figure 2. Reference-guided anomaly detection by feature type. The codon-order shuffle condition demonstrates the main usecase: composition features fail, while sonification features remain highly informative.")
    add_figure(doc, figs["track"], "Figure 3. Example anomaly score track generated from codon64 sonification distances. The same values were exported as a MIDI track so high-scoring anomaly windows become audible as higher/louder notes.")

    h(doc, "Discussion")
    p(doc, "The revised study supports a focused claim: sonification can represent certain genomic anomaly patterns in a way that ordinary composition summaries do not. The strongest evidence is the codon-order shuffle experiment. Because that anomaly preserves codon composition, a base/codon frequency baseline has little information to use. Sonification features based on pitch intervals and pitch order still detected those anomaly windows with high AUC.")
    p(doc, "This does not mean that sonification discovers biological truth by itself. The analysis was reference-guided, and the anomaly labels were synthetic. The method should be understood as a display and feature-extraction layer that can sit on top of standard sequence comparison. Its value is not replacing BLAST, alignment, or variant calling, but helping represent where a genome differs from a reference and making order-based changes more perceptible.")
    p(doc, "The host-group test was intentionally secondary. Codon64 audio features matched the composition baseline at 40% accuracy, which was above the 20% chance level but far from sufficient for classification. This result is useful mainly because it shows that the audio representation did not erase all biological structure. It should not be used to claim host prediction or pathogenicity prediction.")
    p(doc, "The perceptual rationale for this result is plausible but not yet proven by listener testing. Human hearing is highly sensitive to changes in pitch contour, repeated interval patterns, and sudden melodic disruptions. Because codon-order shuffles preserve the same codon inventory while changing the order in which codons occur, they can leave composition summaries unchanged but alter the interval sequence of a sonified track. This makes auditory display a natural candidate for representing order disruptions, especially as a companion to visual plots and standard sequence comparison.")
    p(doc, "Several limitations remain. First, the anomalies were simulated, not verified natural mutations. Second, the first 60 windows were used for standardization and runtime control, not whole-genome alignment. Third, some host groups required chromosome or scaffold assemblies because complete genomes were not equally available across groups. Fourth, the codon64 mapping is transparent but arbitrary; future work should compare multiple biologically informed pitch maps. Finally, listener studies would be needed to test whether humans can reliably hear the anomaly tracks without seeing the computed scores.")

    h(doc, "Conclusion")
    p(doc, "This paper now shows a concrete usecase for genomic sonification. Across 50 X. fastidiosa RefSeq assemblies, codon-level sonification preserved modest host-associated sequence signal and strongly detected reference-guided order anomalies. The key finding is that codon-order shuffles were invisible to composition features but visible to sonification-derived pitch and interval features. Sonification is therefore best framed as an interpretable anomaly-representation layer for genomics: not a replacement for standard bioinformatics, but a novel way to encode and communicate changes in sequence order.")

    h(doc, "Data and Code Availability")
    p(doc, "All scripts and generated files are stored in the local project workspace. The main commands are python analysis/select_and_download.py and python analysis/run_benchmark.py. The selected genome metadata are in data/study_genomes.csv, and the principal results are in results/summary.json and results/anomaly_detection_summary.csv.")

    h(doc, "References")
    for i, ref in enumerate(REFERENCES, 1):
        p(doc, f"{i}. {ref}")

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    h(doc, "Appendix")
    p(doc, "Appendix Table A1. RefSeq assemblies included in the 50-assembly study dataset.")
    add_accession_table(doc, genomes)

    path = DOCX_DIR / "Hearing_Genomic_Order_Revised_Manuscript_Reviewer_Edits.docx"
    doc.save(path)
    return path


def build_response() -> pathlib.Path:
    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    apply_styles(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Response to Reviewer Feedback")
    set_run_font(run, size=14, bold=True)
    p(doc, "Dear Editor and Reviewers,")
    p(doc, "Thank you for the detailed critique. I treated the review as a request for a full redesign. The manuscript no longer claims that sonification diagnoses pathogenicity or replaces standard bioinformatics. Instead, it tests a narrower and more defensible question: whether deterministic windowed sonification can represent mutation/anomaly patterns, especially order changes that composition summaries miss.")
    h(doc, "Major Revisions")
    for item in [
        "Replaced the earlier three-sequence design with 50 NCBI RefSeq Xylella fastidiosa assemblies, balanced across five host-associated groups.",
        "Corrected the biological framing of X. fastidiosa as a bacterial plant pathogen.",
        "Changed the study aim from broad genome comparison to reference-guided genomic anomaly representation.",
        "Added explicit preprocessing, sonification maps, baseline composition features, permutation testing, AUC, precision@4, bootstrap confidence intervals, and public-data ethics statement.",
        "Added a central result showing a sonification-specific usecase: codon-order shuffle anomalies were missed by composition features but detected by codon64 and amino-class sonification features.",
        "Added figures, full accession appendix, code/data availability, and a revised reference list with the requested sonification literature.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    h(doc, "Remaining Scope Limits")
    p(doc, "The revised manuscript is careful not to claim disease diagnosis, virulence prediction, or replacement of alignment/variant-calling workflows. It presents sonification as an interpretable representation layer for reference-guided anomaly windows.")
    path = DOCX_DIR / "Response_to_Reviewer_Feedback_Revised.docx"
    doc.save(path)
    return path


def main() -> None:
    print(build_manuscript())
    print(build_response())


if __name__ == "__main__":
    main()
