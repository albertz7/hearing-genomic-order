from __future__ import annotations

import pathlib
import re
import shutil

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

import build_final_paper


ROOT = pathlib.Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "output" / "submission"
TITLE_FILENAME = (
    "Hearing Genomic Order - Windowed Sonification for Representing Mutation "
    "and Anomaly Patterns in Xylella fastidiosa Assemblies"
)


def set_font(run, size: int = 12, superscript: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.superscript = superscript


def citation_numbers(token: str) -> list[int]:
    out: list[int] = []
    for part in token.split(","):
        part = part.strip().replace("–", "-")
        match = re.match(r"^(\d+)\s*-\s*(\d+)$", part)
        if match:
            out.extend(range(int(match.group(1)), int(match.group(2)) + 1))
        elif part.isdigit():
            out.append(int(part))
    return out


def clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def add_online_citation(paragraph, numbers: list[int], refs: list[str]) -> None:
    for idx, number in enumerate(numbers):
        run = paragraph.add_run(f" (({refs[number - 1]}))")
        set_font(run)
        if idx != len(numbers) - 1:
            comma = paragraph.add_run(",")
            set_font(comma, superscript=True)


def convert_paragraph_citations(paragraph, refs: list[str]) -> None:
    text = paragraph.text
    matches = list(re.finditer(r"\[([0-9,\-– ]+)\]", text))
    if not matches:
        return
    clear_paragraph(paragraph)
    cursor = 0
    for match in matches:
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_font(run)
        nums = citation_numbers(match.group(1))
        add_online_citation(paragraph, nums, refs)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_font(run)


def prepare_online_docx(standard_path: pathlib.Path, online_path: pathlib.Path) -> None:
    doc = Document(standard_path)
    refs = build_final_paper.REFERENCES
    for paragraph in doc.paragraphs:
        convert_paragraph_citations(paragraph, refs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    convert_paragraph_citations(paragraph, refs)
    doc.save(online_path)


def main() -> None:
    SUBMISSION.mkdir(parents=True, exist_ok=True)
    generated = build_final_paper.build_manuscript()
    standard = SUBMISSION / f"{TITLE_FILENAME}.docx"
    online = SUBMISSION / f"{TITLE_FILENAME} - Online Citation Format.docx"
    shutil.copyfile(generated, standard)
    prepare_online_docx(standard, online)
    print(standard)
    print(online)


if __name__ == "__main__":
    main()
